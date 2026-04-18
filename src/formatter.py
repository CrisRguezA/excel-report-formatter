"""
formatter.py

ES:
Genera un reporte Excel profesional a partir de datos de ventas limpios.
Produce tres hojas: KPIs (métricas clave), Pivot (tablas cruzadas) y
Charts (gráficos de barras y tarta).

Contrato del pipeline:
    generate_report(input_path: str, output_path: str) -> None
    input_path  : ruta al Excel limpio (output del excel-data-cleaner)
    output_path : ruta donde se guarda el reporte final

EN:
Generates a professional Excel report from clean sales data.
Produces three sheets: KPIs (key metrics), Pivot (cross tables) and
Charts (bar and pie charts).
"""

# ------------------------------------------------------------------
# IMPORTS
# ------------------------------------------------------------------

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.series import DataPoint


# ------------------------------------------------------------------
# STYLE CONSTANTS  (reutilizados de formatting.py del cleaner)
# ------------------------------------------------------------------

HEADER_FILL_BLUE  = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
HEADER_FILL_GREEN = PatternFill("solid", start_color="375623", end_color="375623")
HEADER_FILL_DARK  = PatternFill("solid", start_color="2E4057", end_color="2E4057")
KPI_LABEL_FILL    = PatternFill("solid", start_color="D6E4F0", end_color="D6E4F0")
KPI_VALUE_FILL    = PatternFill("solid", start_color="EBF3FB", end_color="EBF3FB")
ALT_FILL          = PatternFill("solid", start_color="EBF3FB", end_color="EBF3FB")

HEADER_FONT  = Font(name="Arial", bold=True, color="FFFFFF", size=11)
TITLE_FONT   = Font(name="Arial", bold=True, color="1F4E79", size=14)
LABEL_FONT   = Font(name="Arial", bold=True, size=10)
BODY_FONT    = Font(name="Arial", size=10)
KPI_VAL_FONT = Font(name="Arial", bold=True, size=12, color="1F4E79")

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left",   vertical="center")


# ------------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------------

def _autofit_columns(ws, min_width=10, max_width=40):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_width), max_width)


def _apply_header_row(ws, fill):
    for cell in ws[1]:
        cell.font  = HEADER_FONT
        cell.fill  = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def _style_table(ws, header_fill):
    """Aplica estilo completo a una tabla: encabezado + cuerpo + autofit."""
    _apply_header_row(ws, header_fill)
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill = ALT_FILL if i % 2 == 0 else PatternFill()
        for cell in row:
            cell.font   = BODY_FONT
            cell.border = THIN_BORDER
            cell.fill   = fill
            cell.alignment = CENTER
    _autofit_columns(ws)
    ws.row_dimensions[1].height = 28


# ------------------------------------------------------------------
# SHEET 1 — KPIs
# ------------------------------------------------------------------

def _build_kpis(ws, df):
    """Construye la hoja KPIs con métricas clave de ventas y certificación."""

    # --- Cálculos ---
    total_ventas      = len(df)
    importe_total     = df["importe"].sum()
    ticket_medio      = df["importe"].mean()
    ventas_cerradas   = (df["estado"] == "Cerrado").sum()
    pct_cerradas      = ventas_cerradas / total_ventas * 100

    certificadas      = df[df["certificacion"] != "Sin certificación"]
    total_cert        = len(certificadas)
    pct_certificadas  = total_cert / total_ventas * 100
    importe_cert      = certificadas["importe"].sum()
    importe_no_cert   = df[df["certificacion"] == "Sin certificación"]["importe"].sum()

    # --- Título ---
    ws.merge_cells("B2:D2")
    title_cell = ws["B2"]
    title_cell.value     = "Reporte de Ventas — KPIs"
    title_cell.font      = TITLE_FONT
    title_cell.alignment = CENTER

    # --- Datos de las tarjetas KPI ---
    kpis = [
        ("Total de ventas",              total_ventas,       "#"),
        ("Importe total",                importe_total,      "€"),
        ("Ticket medio",                 ticket_medio,       "€"),
        ("Ventas cerradas",              ventas_cerradas,    "#"),
        ("% ventas cerradas",            pct_cerradas,       "%"),
        ("Ventas certificadas",          total_cert,         "#"),
        ("% ventas certificadas",        pct_certificadas,   "%"),
        ("Importe certificado",          importe_cert,       "€"),
        ("Importe sin certificación",    importe_no_cert,    "€"),
    ]

    # --- Escribir tarjetas (fila inicio: 4) ---
    ws.column_dimensions["A"].width = 3   # margen izquierdo
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 8

    start_row = 4
    for i, (label, value, unit) in enumerate(kpis):
        row = start_row + i

        # Etiqueta
        label_cell = ws.cell(row=row, column=2, value=label)
        label_cell.font      = LABEL_FONT
        label_cell.fill      = KPI_LABEL_FILL
        label_cell.border    = THIN_BORDER
        label_cell.alignment = LEFT

        # Valor formateado
        if unit == "€":
            display = f"{value:,.2f} €"
        elif unit == "%":
            display = f"{value:.1f} %"
        else:
            display = str(int(value))

        value_cell = ws.cell(row=row, column=3, value=display)
        value_cell.font      = KPI_VAL_FONT
        value_cell.fill      = KPI_VALUE_FILL
        value_cell.border    = THIN_BORDER
        value_cell.alignment = CENTER

        ws.row_dimensions[row].height = 22

    ws.sheet_view.showGridLines = False


# ------------------------------------------------------------------
# SHEET 2 — Pivot
# ------------------------------------------------------------------

def _build_pivot(ws, df):
    """Construye dos tablas pivot: comercial×producto y certificacion×producto."""

    # --- Pivot 1: Importe por Comercial × Producto ---
    pivot1 = df.pivot_table(
        index="comercial",
        columns="producto",
        values="importe",
        aggfunc="sum",
        fill_value=0
    ).reset_index()
    pivot1.columns.name = None

    # Título
    ws["A1"] = "Importe por Comercial y Producto (€)"
    ws["A1"].font      = TITLE_FONT
    ws["A1"].alignment = LEFT
    ws.row_dimensions[1].height = 28

    # Encabezados
    headers1 = list(pivot1.columns)
    for col_idx, header in enumerate(headers1, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL_BLUE
        cell.alignment = CENTER
        cell.border    = THIN_BORDER

    # Datos pivot1
    for row_idx, row_data in enumerate(pivot1.itertuples(index=False), start=3):
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font      = BODY_FONT
            cell.border    = THIN_BORDER
            cell.alignment = CENTER
            if row_idx % 2 == 0:
                cell.fill = ALT_FILL

    # --- Pivot 2: Importe por Certificación × Producto ---
    pivot2 = df.pivot_table(
        index="certificacion",
        columns="producto",
        values="importe",
        aggfunc="sum",
        fill_value=0
    ).reset_index()
    pivot2.columns.name = None

    offset_row = len(pivot1) + 5  # separación entre tablas

    ws.cell(row=offset_row, column=1).value     = "Importe por Certificación y Producto (€)"
    ws.cell(row=offset_row, column=1).font      = TITLE_FONT
    ws.cell(row=offset_row, column=1).alignment = LEFT
    ws.row_dimensions[offset_row].height = 28

    headers2 = list(pivot2.columns)
    for col_idx, header in enumerate(headers2, start=1):
        cell = ws.cell(row=offset_row + 1, column=col_idx, value=header)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL_GREEN
        cell.alignment = CENTER
        cell.border    = THIN_BORDER

    for row_idx, row_data in enumerate(pivot2.itertuples(index=False), start=offset_row + 2):
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font      = BODY_FONT
            cell.border    = THIN_BORDER
            cell.alignment = CENTER
            if row_idx % 2 == 0:
                cell.fill = ALT_FILL

    _autofit_columns(ws)


# ------------------------------------------------------------------
# SHEET 3 — Charts
# ------------------------------------------------------------------

def _build_charts(ws, df):
    """Construye 4 gráficos: barras por producto, tarta estado,
    tarta certificación, tarta tipo madera."""

    # --- Datos auxiliares escritos en la hoja (base para los gráficos) ---

    # 1. Importe por producto
    prod_data = df.groupby("producto")["importe"].sum().reset_index()
    ws["A1"] = "Producto"
    ws["B1"] = "Importe"
    for i, row in prod_data.iterrows():
        ws.cell(row=i+2, column=1, value=row["producto"])
        ws.cell(row=i+2, column=2, value=row["importe"])

    # 2. Ventas por estado
    est_data = df["estado"].value_counts().reset_index()
    est_data.columns = ["estado", "count"]
    ws["D1"] = "Estado"
    ws["E1"] = "Ventas"
    for i, row in est_data.iterrows():
        ws.cell(row=i+2, column=4, value=row["estado"])
        ws.cell(row=i+2, column=5, value=row["count"])

    # 3. Ventas por certificación
    cert_data = df["certificacion"].value_counts().reset_index()
    cert_data.columns = ["certificacion", "count"]
    ws["G1"] = "Certificación"
    ws["H1"] = "Ventas"
    for i, row in cert_data.iterrows():
        ws.cell(row=i+2, column=7, value=row["certificacion"])
        ws.cell(row=i+2, column=8, value=row["count"])

    # 4. Ventas por tipo madera
    mad_data = df["tipo_madera"].value_counts().reset_index()
    mad_data.columns = ["tipo_madera", "count"]
    ws["J1"] = "Tipo Madera"
    ws["K1"] = "Ventas"
    for i, row in mad_data.iterrows():
        ws.cell(row=i+2, column=10, value=row["tipo_madera"])
        ws.cell(row=i+2, column=11, value=row["count"])

    n_prod = len(prod_data)
    n_est  = len(est_data)
    n_cert = len(cert_data)
    n_mad  = len(mad_data)

    # --- Gráfico 1: Barras — Importe por Producto ---
    bar = BarChart()
    bar.type    = "col"
    bar.title   = "Importe por Producto (€)"
    bar.y_axis.title = "Importe (€)"
    bar.x_axis.title = "Producto"
    bar.style   = 10
    bar.width   = 14
    bar.height  = 10

    data_ref = Reference(ws, min_col=2, min_row=1, max_row=n_prod+1)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=n_prod+1)
    bar.add_data(data_ref, titles_from_data=True)
    bar.set_categories(cats_ref)
    ws.add_chart(bar, "A10")

    # --- Gráfico 2: Tarta — Ventas por Estado ---
    pie_est = PieChart()
    pie_est.title  = "Ventas por Estado"
    pie_est.style  = 10
    pie_est.width  = 12
    pie_est.height = 10

    data_ref2 = Reference(ws, min_col=5, min_row=1, max_row=n_est+1)
    cats_ref2 = Reference(ws, min_col=4, min_row=2, max_row=n_est+1)
    pie_est.add_data(data_ref2, titles_from_data=True)
    pie_est.set_categories(cats_ref2)
    ws.add_chart(pie_est, "F10")

    # --- Gráfico 3: Tarta — Ventas por Certificación ---
    pie_cert = PieChart()
    pie_cert.title  = "Ventas por Certificación"
    pie_cert.style  = 10
    pie_cert.width  = 12
    pie_cert.height = 10

    data_ref3 = Reference(ws, min_col=8, min_row=1, max_row=n_cert+1)
    cats_ref3 = Reference(ws, min_col=7, min_row=2, max_row=n_cert+1)
    pie_cert.add_data(data_ref3, titles_from_data=True)
    pie_cert.set_categories(cats_ref3)
    ws.add_chart(pie_cert, "A28")

    # --- Gráfico 4: Tarta — Ventas por Tipo de Madera ---
    pie_mad = PieChart()
    pie_mad.title  = "Ventas por Tipo de Madera"
    pie_mad.style  = 10
    pie_mad.width  = 12
    pie_mad.height = 10

    data_ref4 = Reference(ws, min_col=11, min_row=1, max_row=n_mad+1)
    cats_ref4 = Reference(ws, min_col=10, min_row=2, max_row=n_mad+1)
    pie_mad.add_data(data_ref4, titles_from_data=True)
    pie_mad.set_categories(cats_ref4)
    ws.add_chart(pie_mad, "F28")

    ws.sheet_view.showGridLines = False


# ------------------------------------------------------------------
# MAIN FUNCTION
# ------------------------------------------------------------------

def generate_report(input_path: str, output_path: str) -> None:
    """
    Punto de entrada principal.
    Lee el Excel limpio, genera el reporte con 3 hojas y lo guarda.

    Args:
        input_path  : ruta al Excel limpio (output del excel-data-cleaner)
        output_path : ruta donde se guarda el reporte final
    """
    df = pd.read_excel(input_path)

    wb = Workbook()

    # Renombrar hoja por defecto y crear las demás
    ws_kpis   = wb.active
    ws_kpis.title = "KPIs"
    ws_pivot  = wb.create_sheet("Pivot")
    ws_charts = wb.create_sheet("Charts")

    _build_kpis(ws_kpis, df)
    _build_pivot(ws_pivot, df)
    _build_charts(ws_charts, df)

    wb.save(output_path)
    print(f"✓ Reporte generado: {output_path}")
